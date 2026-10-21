"""Read the authoritative XLSX, including cached values and formula provenance. Never save it."""
import datetime
import hashlib
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/raw/KelompokB_MockData_Software_NUSA_CerviScan_cobas.xlsx"
OUTPUT = ROOT / "data/generated/workbook.json"
TABLES = {
    "dictionary": ("1 Kamus Data", 1),
    "nusa": ("2 Mock NUSA", 1),
    "cerviscan": ("3 Mock CerviScan", 2),
    "cobas": ("4 Mock cobas", 2),
    "canonical": ("5 Format Standar", 2),
    "followups": ("6 Tindak Lanjut", 2),
    "summary": ("7 Ringkasan", 3),
}


def serial(value):
    if isinstance(value, datetime.datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, datetime.date):
        return value.isoformat()
    return value


def extract():
    before = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    values = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(SOURCE, read_only=True, data_only=False)
    output = {
        "metadata": {
            "sourceFile": SOURCE.relative_to(ROOT).as_posix(),
            "sha256": before,
            "dataVersion": "2026-10-19",
            "synthetic": True,
            "mode": "recorded-scenario-snapshot",
            "sheets": [],
        },
        "assumptions": {},
        "notes": [],
        "tables": {},
    }
    for sheet in values:
        formula_sheet = formulas[sheet.title]
        output["metadata"]["sheets"].append({
            "name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
            "formulaCount": sum(c.data_type == "f" for r in formula_sheet for c in r),
        })
    for key, (name, header_row) in TABLES.items():
        cached_rows = list(values[name].iter_rows())
        formula_rows = list(formulas[name].iter_rows())
        headers = [c.value for c in cached_rows[header_row - 1]]
        if any(not isinstance(h, str) for h in headers) or len(headers) != len(set(headers)):
            raise ValueError(f"Invalid table headers: {name}")
        records = []
        for cached, original in zip(cached_rows[header_row:], formula_rows[header_row:]):
            if not any(c.value is not None for c in cached):
                continue
            # The final summary explanatory note is not an indicator record.
            if key == "summary" and cached[0].row > 17:
                continue
            records.append({
                "excelRow": cached[0].row,
                "values": {h: serial(c.value) for h, c in zip(headers, cached)},
                "formulas": {c.coordinate: c.value for c in original if c.data_type == "f"},
            })
        output["tables"][key] = {"sheet": name, "headerRow": header_row, "headers": headers, "rows": records}
    config = {
        "hpvCtPositiveThreshold": 13, "internalControlMaxCt": 14,
        "branchBShippingWindowDays": 15, "ivaTargetDays": 16,
        "negativeRecallYears": 17, "otherNormalIvaRecallYears": 18,
        "positive1618NormalIvaRecallYears": 19,
    }
    for key, row in config.items():
        output["assumptions"][key] = {
            "value": values["Baca Dulu"].cell(row, 2).value,
            "source": f"Baca Dulu!B{row}",
            "note": values["Baca Dulu"].cell(row, 3).value,
            "clinicalValidated": False,
        }
    output["notes"] = [[serial(c.value) for c in r] for r in values["Baca Dulu"]]
    values.close()
    formulas.close()
    if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != before:
        raise RuntimeError("Source workbook changed during extraction")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Extracted {len(output['tables']['canonical']['rows'])} reference records; workbook unchanged ({before}).")


if __name__ == "__main__":
    extract()
